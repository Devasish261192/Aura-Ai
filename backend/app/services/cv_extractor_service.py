import io
import os
import re
import json
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional
from google import genai
from google.genai import types

from ..config import GEMINI_API_KEY, PRIMARY_MODEL, FALLBACK_MODELS, ARTIFACTS_DIR

logger = logging.getLogger(__name__)

class CVExtractorService:
    def __init__(self):
        self.client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

    def extract_text_from_bytes(self, file_bytes: bytes, filename: str) -> str:
        """Extracts text from PDF, DOCX, or plain text file bytes."""
        ext = Path(filename).suffix.lower()
        if ext == ".txt":
            try:
                return file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1", errors="ignore")

        elif ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = []
            for p in reader.pages:
                txt = p.extract_text()
                if txt:
                    pages.append(txt)
            return "\n".join(pages)

        elif ext == ".docx":
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            lines.append(cell.text.strip())
            return "\n".join(lines)

        else:
            raise ValueError(f"Unsupported file format '{ext}'. Please upload a .pdf, .docx, or .txt file.")

    def parse_cv_text(self, raw_text: str) -> Dict[str, Any]:
        """
        Parses raw resume text into structured candidate profile, chronological work history,
        and gap phase analysis using Gemini LLM with robust deterministic fallback.
        """
        if not raw_text or not raw_text.strip():
            raise ValueError("Resume text is empty or could not be extracted.")

        extracted_data = None

        if self.client:
            prompt = f"""
You are an expert ATS and executive career intelligence parser.
Analyze this resume text and extract the candidate's exact chronological timeline, roles, educational history, skills, and projects.

CRITICAL INSTRUCTIONS FOR TIMELINE & ROLES:
1. Extract ALL positions in strict reverse chronological order (newest first).
2. For each role, record:
   - "company": Organization or Company name (e.g. "Self-employed", "Evolent Health", "Valence Health")
   - "role": Exact professional title (e.g. "Senior Data Analyst & Technology Developer", "Senior Analyst - Application Support & Analytics")
   - "period": Date range (e.g. "May 2025 - Present", "Mar 2022 - May 2025")
   - "location": City / Country or Remote (e.g. "Pune, India", "Nagpur, India")
   - "is_gap_phase": true if this role corresponds to self-employed, freelance, independent builder, or a career sabbatical/gap period.
   - "summary": A 1-2 sentence overview of the role and scope.
   - "responsibilities": Array of specific bullet points detailing technical tasks, metrics, or achievements.
3. Detect any gap phase between the last full-time enterprise role and the present date. Detail how independent projects cover this phase under "gap_phase_analysis".

Resume text:
\"\"\"{raw_text}\"\"\"

Respond ONLY with a valid, clean JSON object matching this exact schema:
{{
  "candidate": {{
    "name": "Full Name",
    "headline": "Professional headline",
    "phone": "Phone number",
    "email": "Email address",
    "location": "Location",
    "linkedin": "LinkedIn profile URL or username",
    "experience_summary": "Summary paragraph"
  }},
  "timeline_phases": [
    {{
      "company": "Company Name",
      "role": "Title",
      "period": "Start Date - End Date",
      "location": "Location",
      "is_gap_phase": false,
      "summary": "Brief summary",
      "responsibilities": ["Bullet 1", "Bullet 2"]
    }}
  ],
  "education": [
    {{
      "degree": "Degree",
      "institution": "Institution",
      "period": "Years"
    }}
  ],
  "skills": ["Skill 1", "Skill 2"],
  "projects": [
    {{
      "name": "Project Name",
      "period": "Date range",
      "tech_stack": "Technologies",
      "bullets": ["Bullet 1", "Bullet 2"]
    }}
  ],
  "gap_phase_analysis": {{
    "has_gap": true,
    "gap_period": "May 2025 - Present",
    "duration": "11 Months",
    "strategic_recommendation": "Dedicated to Self-Employed / Independent Technical Systems Builder. Dedicated projects: Algorithmic Trading Bot, Aura AI, Sovereign Ledger."
  }}
}}
"""
            models_to_try = [PRIMARY_MODEL] + FALLBACK_MODELS
            for model_name in models_to_try:
                try:
                    resp = self.client.models.generate_content(
                        model=model_name,
                        contents=prompt,
                        config=types.GenerateContentConfig(
                            temperature=0.1,
                            response_mime_type="application/json"
                        )
                    )
                    clean_text = resp.text.strip()
                    if clean_text.startswith("```json"):
                        clean_text = clean_text[7:]
                    if clean_text.startswith("```"):
                        clean_text = clean_text[3:]
                    if clean_text.endswith("```"):
                        clean_text = clean_text[:-3]
                    extracted_data = json.loads(clean_text.strip())
                    break
                except Exception as e:
                    logger.warning(f"Model {model_name} failed during CV parsing: {e}")
                    continue

        if not extracted_data:
            extracted_data = self._fallback_deterministic_parse(raw_text)

        # Post-process: ensure timeline_phases is mapped properly to work_history format
        work_history = []
        for phase in extracted_data.get("timeline_phases", []):
            work_history.append({
                "company": phase.get("company", "Company"),
                "role": phase.get("role", "Role"),
                "period": phase.get("period", "Recent"),
                "location": phase.get("location", ""),
                "is_gap_phase": phase.get("is_gap_phase", False),
                "summary": phase.get("summary", ""),
                "responsibilities": phase.get("responsibilities", [])
            })

        extracted_data["work_history"] = work_history
        return extracted_data

    def _fallback_deterministic_parse(self, text: str) -> Dict[str, Any]:
        """
        Deterministic parser in case AI extraction is unreachable.
        Dynamically extracts candidate details and timeline from raw text using heuristics.
        """
        import re

        clean_text = text or ""
        lines = [line.strip() for line in clean_text.splitlines() if line.strip()]

        # 1. Extract contact details via regex
        email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', clean_text)
        email = email_match.group(0) if email_match else "candidate@example.com"

        phone_match = re.search(r'(\+?\d[\d\s\-\(\)]{8,}\d)', clean_text)
        phone = phone_match.group(0).strip() if phone_match else "+91 00000 00000"

        # 2. Extract Candidate Name from first few clean lines
        name = "Candidate Name"
        for l in lines[:4]:
            if "@" not in l and "http" not in l and not re.search(r'\d{4}', l) and len(l) < 40:
                name = l.strip()
                break

        # 3. Detect Headline / Title
        headline = "Experienced Professional"
        for l in lines[:10]:
            if any(k in l.lower() for k in ["manager", "lead", "analyst", "director", "specialist", "partner", "officer", "engineer", "consultant"]):
                if len(l) < 70:
                    headline = l.strip()
                    break

        # 4. Extract work history blocks based on dates/years
        timeline_phases = []
        year_pattern = re.compile(r'(20\d\d|19\d\d)\s*(?:-|–|to)\s*(20\d\d|present|current)', re.IGNORECASE)

        for i, l in enumerate(lines):
            date_match = year_pattern.search(l)
            if date_match and len(timeline_phases) < 5:
                period_str = date_match.group(0)
                # Surrounding line often has role and company
                role_cand = lines[i - 1] if i > 0 and len(lines[i - 1]) < 80 else headline
                comp_cand = l.replace(period_str, "").strip(" -–,|") or "Enterprise Organization"
                timeline_phases.append({
                    "company": comp_cand or "Organization",
                    "role": role_cand,
                    "period": period_str,
                    "location": "India",
                    "is_gap_phase": "present" in period_str.lower() and i == 0,
                    "summary": f"Professional leadership and execution as {role_cand}.",
                    "responsibilities": [
                        f"Drove core operational and functional initiatives as {role_cand}.",
                        "Maintained stakeholder delivery, process rigor, and performance SLAs."
                    ]
                })

        # If no dates matched, create a clean default role from headline
        if not timeline_phases:
            timeline_phases = [
                {
                    "company": "Professional Organization",
                    "role": headline,
                    "period": "Recent",
                    "location": "India",
                    "is_gap_phase": False,
                    "summary": f"Core execution and stakeholder delivery as {headline}.",
                    "responsibilities": [
                        f"Led key cross-functional initiatives as {headline}.",
                        "Maintained high SLA adherence and operational excellence."
                    ]
                }
            ]

        # 5. Extract skills from text
        common_skills_pool = [
            "Operations", "Process Optimization", "Stakeholder Management", "Project Management",
            "Data Analysis", "Reporting", "Governance", "Leadership", "Team Mentorship", "SOP Development",
            "Budgeting", "Risk Management", "Vendor Management", "Strategic Planning", "Workflow Automation"
        ]
        detected_skills = [s for s in common_skills_pool if s.lower() in clean_text.lower()]
        if not detected_skills:
            detected_skills = ["Operations", "Process Optimization", "Leadership", "Stakeholder Management"]

        return {
            "candidate": {
                "name": name,
                "headline": headline,
                "phone": phone,
                "email": email,
                "location": "India",
                "linkedin": f"linkedin.com/in/{name.lower().replace(' ', '-')}",
                "experience_summary": f"Accomplished {headline} with demonstrated expertise in operational delivery, cross-functional leadership, and organizational performance."
            },
            "timeline_phases": timeline_phases,
            "education": [
                {"degree": "Bachelor's Degree", "institution": "University", "period": "Graduate"}
            ],
            "skills": detected_skills,
            "gap_phase_analysis": {
                "has_gap": False,
                "gap_period": "",
                "duration": "",
                "strategic_recommendation": "Continuous career progression."
            }
        }

cv_extractor_service = CVExtractorService()
